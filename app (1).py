import io
import json
import os
import re
from typing import Any

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


MODEL_NAME = "gemini-2.5-flash"


def get_api_key() -> str:
    """Read the Gemini API key from Streamlit secrets or an environment variable."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "")


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Include text from tables, which are commonly used in resumes.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs).strip()


def extract_text(uploaded_file) -> str:
    data = uploaded_file.getvalue()
    suffix = uploaded_file.name.lower().rsplit(".", 1)[-1]

    if suffix == "pdf":
        return extract_pdf_text(data)
    if suffix == "docx":
        return extract_docx_text(data)
    if suffix in {"txt", "md"}:
        return data.decode("utf-8", errors="replace").strip()

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, TXT, or MD.")


def heuristic_checks(resume_text: str) -> dict[str, Any]:
    """Fast local checks that complement the AI review."""
    text = resume_text.lower()
    word_count = len(re.findall(r"\b[\w+.#/-]+\b", resume_text))
    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", resume_text))
    has_phone = bool(re.search(r"(\+?\d[\d\s().-]{7,}\d)", resume_text))
    common_sections = {
        "experience": bool(re.search(r"\b(work experience|experience|employment)\b", text)),
        "education": bool(re.search(r"\beducation\b", text)),
        "skills": bool(re.search(r"\b(skills|technical skills|core competencies)\b", text)),
        "summary": bool(re.search(r"\b(summary|professional summary|profile)\b", text)),
    }

    return {
        "word_count": word_count,
        "has_email": has_email,
        "has_phone": has_phone,
        "sections": common_sections,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict[str, Any]:
    client = genai.Client(api_key=api_key)

    schema = {
        "type": "OBJECT",
        "properties": {
            "ats_score": {"type": "INTEGER", "minimum": 0, "maximum": 100},
            "verdict": {"type": "STRING"},
            "summary": {"type": "STRING"},
            "strengths": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "improvements": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "missing_keywords": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "formatting_risks": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "action_plan": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
        },
        "required": [
            "ats_score",
            "verdict",
            "summary",
            "strengths",
            "improvements",
            "missing_keywords",
            "formatting_risks",
            "action_plan",
        ],
    }

    job_context = job_description.strip() or (
        "No job description was provided. Evaluate ATS compatibility against "
        "general resume/ATS best practices rather than a specific job."
    )

    prompt = f"""
You are an expert ATS resume reviewer.

Evaluate the resume below. Give an ATS compatibility score from 0 to 100.
If a job description is provided, prioritize keyword alignment with that job.
If no job description is provided, score general ATS readability, structure,
keyword usefulness, measurable achievements, and formatting.

Important:
- Do not invent experience, qualifications, employers, degrees, dates, or skills.
- Missing keywords means useful job-relevant terms that are absent or weak.
- Formatting risks should focus on things that can hurt ATS parsing.
- Make recommendations concrete and actionable.
- Keep the response concise enough for a Streamlit dashboard.

JOB DESCRIPTION:
{job_context}

RESUME:
{resume_text}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=schema,
            temperature=0.2,
            max_output_tokens=3000,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    result = json.loads(response.text)

    # Defensive normalization.
    result["ats_score"] = max(0, min(100, int(result.get("ats_score", 0))))
    for key in (
        "strengths",
        "improvements",
        "missing_keywords",
        "formatting_risks",
        "action_plan",
    ):
        if not isinstance(result.get(key), list):
            result[key] = []

    return result


def render_bullets(items: list[str]) -> None:
    if items:
        for item in items:
            st.markdown(f"- {item}")
    else:
        st.write("None identified.")


st.set_page_config(
    page_title="Resume ATS Checker",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Checker")
st.caption("Upload your resume and get an AI-powered ATS score, risks, and improvement plan.")

with st.sidebar:
    st.header("Settings")
    st.info(
        "This app uses Gemini 2.5 Flash. Your API key is read from "
        "Streamlit Secrets and is never displayed in the app."
    )
    st.markdown(
        "Supported files: **PDF, DOCX, TXT, MD**"
    )

api_key = get_api_key()
if not api_key:
    st.warning(
        "Gemini API key is not configured. Add `GEMINI_API_KEY` in "
        "Streamlit Secrets before running an analysis."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt", "md"],
    help="For best results, upload a text-based PDF or DOCX.",
)

job_description = st.text_area(
    "Optional: paste the job description",
    height=180,
    placeholder="Paste the job description here for a job-specific ATS score and keyword analysis.",
)

analyze_clicked = st.button(
    "Analyze Resume",
    type="primary",
    use_container_width=True,
    disabled=not (uploaded_file and api_key),
)

if analyze_clicked:
    try:
        with st.spinner("Reading and analyzing your resume..."):
            resume_text = extract_text(uploaded_file)

            if len(resume_text.strip()) < 100:
                st.error(
                    "Very little text could be extracted. If this is a scanned/image-only PDF, "
                    "please upload a text-based PDF or DOCX."
                )
                st.stop()

            checks = heuristic_checks(resume_text)
            result = analyze_resume(resume_text, job_description, api_key)

        score = result["ats_score"]

        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("ATS Score", f"{score}/100")
        with col2:
            st.metric("Resume Words", checks["word_count"])
        with col3:
            contact_ok = checks["has_email"] and checks["has_phone"]
            st.metric("Contact Info", "Found" if contact_ok else "Check")

        st.progress(score / 100)
        st.subheader(result.get("verdict", "ATS Review"))
        st.write(result.get("summary", ""))

        tab1, tab2, tab3, tab4 = st.tabs(
            ["Strengths", "Improvements", "Keywords", "Formatting Risks"]
        )

        with tab1:
            render_bullets(result["strengths"])

        with tab2:
            render_bullets(result["improvements"])

        with tab3:
            if job_description.strip():
                render_bullets(result["missing_keywords"])
            else:
                st.info("Add a job description for targeted missing-keyword analysis.")

        with tab4:
            render_bullets(result["formatting_risks"])

        st.subheader("Action Plan")
        render_bullets(result["action_plan"])

        st.subheader("Quick local checks")
        section_status = checks["sections"]
        for section, found in section_status.items():
            st.write(f"{'✅' if found else '⚠️'} {section.title()}: {'found' if found else 'not detected'}")

        with st.expander("View extracted resume text"):
            st.text(resume_text)

    except Exception as exc:
        st.error(f"Analysis failed: {exc}")
        st.caption(
            "Check that your Gemini API key is valid and that the uploaded file "
            "contains selectable text."
        )

st.divider()
st.caption("Privacy note: the app does not save uploaded resumes to disk or a database.")
